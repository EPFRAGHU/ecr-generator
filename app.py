"""
ECR Register & Generator — Flask backend.

Reproduces the logic of ECRfile_Template_v2_1.xlsm (Sheet1 formulas +
TEXTFILE / TEXTSUB macros) as HTTP endpoints, so the browser-side ledger
can generate the same #~#-delimited ECR / Exit text files EPFO expects.

Wage data is stateless — one session's worth per request, same as the
original Excel macro (open file, fill sheet, click button, get a .txt
file). Visitors register (name, mobile, email) before they can generate
files; registrations and usage counts live in SQLite (see visitors.py)
and are viewed at /admin.
"""
from __future__ import annotations

import calendar
import csv
import hashlib
import hmac
import io
import math
import os
import re
import secrets
import time
from datetime import datetime, timedelta, timezone

from flask import (
    Flask,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash

import visitors

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024  # 8MB upload cap
# Behind Coolify's proxy: trust one hop for client IP / scheme.
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)
visitors.init_db()

IST = timezone(timedelta(hours=5, minutes=30))

# Release shown under "Employer Module". Bump on every release:
# MAJOR.MINOR for a feature release, the last number (sub-version) for a
# fix or small change; set APP_BUILD to the actual date and time (IST)
# of the release commit, read from the clock (never estimated).
# History is in README.md under "Version history".
APP_VERSION = "1.10.0"
APP_BUILD = "19-10-2026 11:28 IST"


@app.context_processor
def inject_version() -> dict:
    return {"app_version": APP_VERSION, "app_build": APP_BUILD}

# ---------------------------------------------------------------------------
# Admin login — protects /admin (the visitor list). Credentials from
# environment (set in Coolify):
#   ECR_USERNAME        user ID
#   ECR_PASSWORD_HASH   werkzeug password hash (preferred), or
#   ECR_PASSWORD        plain password
#   SECRET_KEY          session signing key (recommended)
# If no credentials are configured, /admin stays locked.
# ---------------------------------------------------------------------------
LOGIN_USERNAME = os.environ.get("ECR_USERNAME", "").strip()
LOGIN_PASSWORD_HASH = os.environ.get("ECR_PASSWORD_HASH", "").strip()
LOGIN_PASSWORD = os.environ.get("ECR_PASSWORD", "")
LOGIN_CONFIGURED = bool(LOGIN_USERNAME and (LOGIN_PASSWORD_HASH or LOGIN_PASSWORD))

app.secret_key = os.environ.get("SECRET_KEY") or hashlib.sha256(
    ("ecr-generator-session:" + LOGIN_PASSWORD_HASH + LOGIN_PASSWORD).encode()
).hexdigest()
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SECURE=True,
    SESSION_COOKIE_SAMESITE="Lax",
    # Long-lived so registered visitors aren't asked again; admin access
    # expires separately after ADMIN_SESSION_SECONDS.
    PERMANENT_SESSION_LIFETIME=timedelta(days=365),
)
ADMIN_SESSION_SECONDS = 8 * 60 * 60

MAX_FAILED_LOGINS = 5
LOCKOUT_SECONDS = 15 * 60
_failed_logins: dict[str, list[float]] = {}  # client IP -> failure timestamps

MAX_REGISTRATIONS_PER_HOUR = 10
_registrations_by_ip: dict[str, list[float]] = {}

ADMIN_ENDPOINTS = {"admin", "admin_export", "admin_delete"}
REGISTERED_ENDPOINTS = {"upload", "import_ecr", "generate_ecr", "generate_exit",
                        "generate_arrear", "upload_arrear", "generate_pohw", "upload_pohw"}


def _recent(log: dict[str, list[float]], ip: str, window: float) -> list[float]:
    cutoff = time.time() - window
    recent = [t for t in log.get(ip, []) if t > cutoff]
    log[ip] = recent
    return recent


def _recent_failures(ip: str) -> list[float]:
    return _recent(_failed_logins, ip, LOCKOUT_SECONDS)


def _is_admin() -> bool:
    return session.get("admin_until", 0) > time.time()


def _current_reg_id() -> int | None:
    reg_id = session.get("reg_id")
    return reg_id if visitors.registration_exists(reg_id) else None


def _visitor_id() -> str:
    if "vid" not in session:
        session.permanent = True
        session["vid"] = secrets.token_hex(8)
    return session["vid"]


def _log(kind: str) -> None:
    visitors.log_event(_visitor_id(), _current_reg_id(), kind)


def _credentials_ok(username: str, password: str) -> bool:
    user_ok = hmac.compare_digest(username.encode(), LOGIN_USERNAME.encode())
    if LOGIN_PASSWORD_HASH:
        pass_ok = check_password_hash(LOGIN_PASSWORD_HASH, password)
    else:
        pass_ok = hmac.compare_digest(password.encode(), LOGIN_PASSWORD.encode())
    return user_ok and pass_ok


def _safe_next(target: str | None) -> str:
    if target and target.startswith("/") and not target.startswith("//"):
        return target
    return url_for("admin")


@app.before_request
def gate():
    if request.endpoint in ADMIN_ENDPOINTS and not _is_admin():
        return redirect(url_for("login", next=request.path))
    if request.endpoint in REGISTERED_ENDPOINTS and _current_reg_id() is None:
        return jsonify(
            {"error": "Please register to use the ECR generator", "needs_registration": True}
        ), 403
    return None


@app.route("/login", methods=["GET", "POST"])
def login():
    if not LOGIN_CONFIGURED:
        return render_template(
            "login.html", error="Admin login is not configured on the server."
        ), 503

    error = None
    if request.method == "POST":
        ip = request.remote_addr or "unknown"
        if len(_recent_failures(ip)) >= MAX_FAILED_LOGINS:
            error = "Too many failed attempts. Try again in 15 minutes."
        elif _credentials_ok(
            request.form.get("username", "").strip(), request.form.get("password", "")
        ):
            _failed_logins.pop(ip, None)
            session.permanent = True
            session["admin_until"] = time.time() + ADMIN_SESSION_SECONDS
            return redirect(_safe_next(request.form.get("next")))
        else:
            _failed_logins.setdefault(ip, []).append(time.time())
            error = "Incorrect user ID or password."
    return render_template(
        "login.html", error=error, next=request.values.get("next", "")
    )


@app.route("/logout", methods=["POST"])
def logout():
    session.pop("admin_until", None)
    return redirect(url_for("login"))


# ---------------------------------------------------------------------------
# Visitor registration
# ---------------------------------------------------------------------------
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _clean_mobile(raw: str) -> str | None:
    digits = re.sub(r"\D", "", raw or "")
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    return digits if re.fullmatch(r"[6-9]\d{9}", digits) else None


@app.route("/api/register", methods=["POST"])
def register():
    data = request.get_json(silent=True) or {}
    if data.get("website"):  # honeypot field, hidden from people
        return jsonify({"error": "Registration failed"}), 400

    name = re.sub(r"\s+", " ", str(data.get("name", ""))).strip()
    mobile = _clean_mobile(str(data.get("mobile", "")))
    email = str(data.get("email", "")).strip().lower()

    errors = {}
    if not 2 <= len(name) <= 80:
        errors["name"] = "Enter your name"
    if mobile is None:
        errors["mobile"] = "Enter a valid 10-digit mobile number"
    if len(email) > 120 or not EMAIL_RE.match(email):
        errors["email"] = "Enter a valid email address"
    if data.get("consent") is not True:
        errors["consent"] = "Please tick the box to continue"
    if errors:
        return jsonify({"error": "Please check the highlighted fields", "fields": errors}), 400

    ip = request.remote_addr or "unknown"
    if len(_recent(_registrations_by_ip, ip, 3600)) >= MAX_REGISTRATIONS_PER_HOUR:
        return jsonify({"error": "Too many registrations from this network. Try later."}), 429
    _registrations_by_ip.setdefault(ip, []).append(time.time())

    reg_id = visitors.register(name, mobile, email)
    session.permanent = True
    session["reg_id"] = reg_id
    _log("register")
    return jsonify({"status": "ok", "name": name, "email": email})


@app.route("/api/signout", methods=["POST"])
def visitor_signout():
    # Forget the registration on this device; the record itself is kept, and
    # registering again with the same mobile number reuses it.
    session.pop("reg_id", None)
    return jsonify({"status": "ok"})


# ---------------------------------------------------------------------------
# Admin: visitor list
# ---------------------------------------------------------------------------
def _ist(ts) -> str:
    if not ts:
        return "—"
    return datetime.fromtimestamp(ts, IST).strftime("%d-%m-%Y %H:%M")


@app.route("/admin")
def admin():
    people = visitors.list_registrations()
    for p in people:
        p["created"] = _ist(p["created_at"])
        p["last"] = _ist(p["last_seen"])
    return render_template("admin.html", stats=visitors.stats(), people=people)


def _csv_safe(value) -> str:
    # Stop spreadsheet apps from treating a cell as a formula.
    text = str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


@app.route("/admin/export.csv")
def admin_export():
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(
        ["Name", "Mobile", "Email", "Registered (IST)", "Last visit (IST)",
         "Page views", "ECR files", "Exit files"]
    )
    for p in visitors.list_registrations():
        writer.writerow([
            _csv_safe(p["name"]), p["mobile"], _csv_safe(p["email"]),
            _ist(p["created_at"]), _ist(p["last_seen"]),
            p["views"], p["ecr_files"], p["exit_files"],
        ])
    buf = io.BytesIO(out.getvalue().encode("utf-8-sig"))  # BOM so Excel reads UTF-8
    stamp = datetime.now(IST).strftime("%Y%m%d")
    return send_file(buf, mimetype="text/csv", as_attachment=True,
                     download_name=f"ecr-generator-visitors-{stamp}.csv")


@app.route("/admin/delete/<int:reg_id>", methods=["POST"])
def admin_delete(reg_id: int):
    visitors.delete_registration(reg_id)
    return redirect(url_for("admin"))

# ---------------------------------------------------------------------------
# Column auto-detection for uploaded Excel/CSV files
# ---------------------------------------------------------------------------
FIELD_SYNONYMS = {
    "uan": ["uan", "universal account number", "pf uan", "member uan"],
    "name": ["member name", "employee name", "name", "emp name"],
    "gross": ["gross wages", "gross wage", "gross salary", "gross"],
    "epf": ["epf wages", "epf wage", "pf wages", "basic wages", "epf basic"],
    "ncp": ["ncp days", "ncp", "lop days", "loss of pay days"],
    "vpf": ["vpf", "vpf amount", "voluntary pf", "voluntary provident fund"],
    "refund": ["refund of advances", "refund", "advance refund", "refund advance"],
}
REQUIRED_FIELDS = ["uan", "name", "gross", "epf"]
OPTIONAL_FIELDS = ["ncp", "vpf", "refund"]


def _normalize_header(h: str) -> str:
    return re.sub(r"\s+", " ", str(h).strip().lower())


def guess_mapping(headers: list[str]) -> dict[str, str | None]:
    normalized = {h: _normalize_header(h) for h in headers}
    mapping: dict[str, str | None] = {}
    for field, synonyms in FIELD_SYNONYMS.items():
        match = None
        for h, norm in normalized.items():
            if norm in synonyms:
                match = h
                break
        if match is None:
            for h, norm in normalized.items():
                if any(s in norm for s in synonyms):
                    match = h
                    break
        mapping[field] = match
    return mapping


# ---------------------------------------------------------------------------
# Core wage calculations (mirrors Sheet1's formulas)
#
# Wage ceiling revised from 15,000 to 25,000 by S.O. 5109(E), effective
# 17.09.2026. September 2026 is a split month (EPFO FAQ Q7/Q9): days 1-16
# at the old ceiling, days 17-30 at the new one, reported in a single ECR.
# ---------------------------------------------------------------------------
OLD_CEILING = 15000
NEW_CEILING = 25000
SPLIT_MONTH = "2026-09"
SPLIT_DAYS_BEFORE, SPLIT_DAYS_AFTER, SPLIT_DAYS_TOTAL = 16, 14, 30

# How a member was treated from 01.09.2026 to 16.09.2026 (FAQ Q7 scenarios).
# Whether EPF is on capped or full wages is the separate "higher EPF" flag.
SEP_BASES = {
    "cap",    # Scenario C: EPS member throughout
    "noeps",  # Scenario B: EPF-only member, joins EPS 17.09.2026
    "new",    # Scenario A: excluded employee, covered from 17.09.2026
}


def round_half_up(x: float) -> int:
    # Python's round() is banker's rounding (2082.5 -> 2082); EPFO rounds up.
    return math.floor(x + 0.5 + 1e-9)


def ceiling_for(wage_month: str) -> int:
    return NEW_CEILING if (wage_month or "") > SPLIT_MONTH else OLD_CEILING


# Higher EPF modes. Keep in sync with HIGHER_MODES in static/app.js.
HIGHER_MODES = {
    "Y",  # employee and employer both on the actual wage above the ceiling
    "E",  # employee on the actual wage, employer restricted to the ceiling
    "N",  # EPF wages capped at the ceiling for both
}


def calc_row(
    epf,
    wage_month="",
    eps_member=True,
    sep_basis="cap",
    higher="N",
    vpf=0,
    higher_pension=False,
) -> dict:
    """
    `epf` is the member's full-month EPF wages. `higher` is a HIGHER_MODES
    value (True/False are read as "Y"/"N"):

      Y  EPF wages are the actual wage; employee and employer both pay 12%
         on it. EPS and EDLI wages stay capped, so the employer's excess
         goes to the EPF diff.
      E  EPF wages are the actual wage and the employee pays 12% on it, but
         the employer's 12% is restricted to the ceiling wage (EPS 8.33% of
         it, the rest as the diff), e.g. 1,250 + 550 at 15,000 and
         2,083 + 917 at 25,000.
      N  EPF wages capped at the ceiling for both (the default).

    The ceiling is 15,000 up to 16.09.2026 and 25,000 from 17.09.2026; in
    September 2026 the wage is split per `sep_basis`.

    `vpf` is the voluntary PF amount in rupees. It is remitted with the
    employee share, so it is added to the EPF contribution but not to the
    employer's diff.

    `higher_pension` is for an EPS member whose joint option for pension on
    higher wages has been accepted. EPF is then on the actual wage for both
    (mode "Y"), EPS wages are the actual wage, and the EPS contribution is
    8.33% of it plus 1.16% of the part above the ceiling, all out of the
    employer's 12%. EDLI wages stay capped.
    """
    if higher is True:
        higher = "Y"
    elif higher is False:
        higher = "N"
    mode = higher if higher in HIGHER_MODES else "N"
    hp = bool(higher_pension) and eps_member
    if hp:
        mode = "Y"

    w = float(epf or 0)
    full_w = w
    if wage_month == SPLIT_MONTH:
        f1 = SPLIT_DAYS_BEFORE / SPLIT_DAYS_TOTAL
        f2 = SPLIT_DAYS_AFTER / SPLIT_DAYS_TOTAL
        old, new = min(w, OLD_CEILING), min(w, NEW_CEILING)
        split = old * f1 + new * f2
        # Wage above the ceiling for each part of the month (higher pension)
        over_old, over_new = max(w - OLD_CEILING, 0), max(w - NEW_CEILING, 0)
        edli_w = split
        excess = 0.0
        if eps_member and sep_basis == "new":
            # Covered only from 17th: wages for 17-30 Sep only.
            full_w = w * f2
            capped_w = new * f2
            eps_w = edli_w = new * f2
            if hp:
                eps_w, excess = w * f2, over_new * f2
        else:
            capped_w = split
            if not eps_member:
                eps_w = 0.0
            elif sep_basis == "noeps":
                eps_w = new * f2
                if hp:
                    eps_w, excess = w * f2, over_new * f2
            else:  # "cap"
                eps_w = split
                if hp:
                    eps_w, excess = w, over_old * f1 + over_new * f2
    else:
        ceiling = ceiling_for(wage_month)
        capped_w = min(w, ceiling)
        eps_w = min(w, ceiling) if eps_member else 0.0
        edli_w = min(w, ceiling)
        excess = 0.0
        if hp:
            eps_w, excess = w, max(w - ceiling, 0)

    epf_w = capped_w if mode == "N" else full_w
    employer_w = full_w if mode == "Y" else capped_w

    ee_contri = round_half_up(epf_w * 0.12)
    er_contri = round_half_up(employer_w * 0.12)
    eps_contri = round_half_up(eps_w * 0.0833 + excess * 0.0116)
    return {
        "epf_wages": round_half_up(epf_w),
        "eps_wages": round_half_up(eps_w),
        "edli_wages": round_half_up(edli_w),
        "epf_contri": ee_contri + round_half_up(float(vpf or 0)),
        "eps_contri": eps_contri,
        "diff": er_contri - eps_contri,
    }


def days_in_month(wage_month: str) -> int:
    try:
        year, month = (int(p) for p in str(wage_month).split("-"))
        return calendar.monthrange(year, month)[1]
    except (TypeError, ValueError):
        return 31


# Keep ncp_issue / vpf_issue in sync with ncpIssue / vpfIssue in static/app.js.
def ncp_issue(ncp, wage_month: str) -> str:
    s = str(ncp if ncp is not None else "").strip()
    if s in ("", "0"):
        return ""
    if not re.fullmatch(r"\d+", s):
        return "NCP days must be whole days"
    max_days = days_in_month(wage_month)
    if int(s) > max_days:
        return f"NCP days cannot exceed {max_days}"
    return ""


def duplicate_uan_issue(rows: list[dict]) -> str:
    """Name the first UAN that appears on more than one row, if any."""
    seen: dict[str, int] = {}
    for n, row in enumerate(rows, start=1):
        uan = str(row.get("uan", "")).strip()
        if not uan:
            continue
        if uan in seen:
            return f"Rows {seen[uan]} and {n} have the same UAN {uan}"
        seen[uan] = n
    return ""


def vpf_issue(vpf) -> str:
    s = str(vpf if vpf is not None else "").strip()
    if s in ("", "0"):
        return ""
    return "" if re.fullmatch(r"\d+(\.\d+)?", s) else "VPF is not a valid amount"


def row_eps_member(row: dict) -> bool:
    return str(row.get("eps", "Y")).strip().upper() != "N"


def row_higher_epf(row: dict) -> str:
    mode = str(row.get("higher", "N")).strip().upper()
    return mode if mode in HIGHER_MODES else "N"


def row_higher_pension(row: dict) -> bool:
    return str(row.get("hp", "N")).strip().upper() == "Y"


def row_sep_basis(row: dict) -> str:
    basis = str(row.get("sep", "cap")).strip().lower()
    return basis if basis in SEP_BASES else "cap"


def validate_wage_row(row: dict) -> list[str]:
    issues = []
    uan = str(row.get("uan", "")).strip()
    if not re.fullmatch(r"\d{12}", uan):
        issues.append("UAN is not a 12-digit number")
    if not str(row.get("name", "")).strip():
        issues.append("Name is blank")
    try:
        gross = float(row.get("gross", 0))
        epf = float(row.get("epf", 0))
        if epf > gross:
            issues.append("EPF wages exceed gross wages")
    except (TypeError, ValueError):
        issues.append("Gross/EPF wages are not numeric")
    # The wage month isn't known at upload time; the page re-checks NCP days
    # against the chosen month.
    for issue in (ncp_issue(row.get("ncp"), ""), vpf_issue(row.get("vpf"))):
        if issue:
            issues.append(issue)
    return issues


def _whole_number_text(val: str) -> str:
    # Excel numbers often arrive as "4.0"; keep them as "4".
    return re.sub(r"^(\d+)\.0+$", r"\1", val)


# ---------------------------------------------------------------------------
# ECR text file import
#
# An ECR line has 11 #~#-separated fields: UAN, name, gross wages, EPF
# wages, EPS wages, EDLI wages, EPF contribution (employee, incl. VPF), EPS
# contribution, EPF diff (employer), NCP days, refund of advances.
#
# The register holds the full-month EPF wage plus the member's options
# (EPS member, higher EPF, higher pension, September 2026 status), so each
# line is worked back to the inputs that make calc_row() reproduce it
# exactly. A line no combination reproduces is still imported, with a note.
# ---------------------------------------------------------------------------
ECR_FIELD_COUNT = 11
ECR_MONTHS = ("JAN", "FEB", "MAR", "APR", "MAY", "JUN",
              "JUL", "AUG", "SEP", "OCT", "NOV", "DEC")
ESTAB_CODE_RE = re.compile(r"[A-Z]{5}\d{10}")


def ecr_filename_hints(filename: str) -> tuple[str | None, str | None]:
    """Establishment code and wage month (YYYY-MM) from an ECR file name,
    e.g. ORBBS2012397000_ECR_SEP_2026_BATCH2.txt or ORBBS0017675000202609.txt."""
    name = (filename or "").upper()
    estab = ESTAB_CODE_RE.search(name)
    month = None
    m = re.search(r"(?<![A-Z])(" + "|".join(ECR_MONTHS) + r")[A-Z]*[_\-\s.]*(20\d{2})(?!\d)", name)
    if m:
        month = f"{m.group(2)}-{ECR_MONTHS.index(m.group(1)) + 1:02d}"
    elif estab:
        m = re.match(r"(20\d{2})(0[1-9]|1[0-2])(?!\d)", name[estab.end():])
        if m:
            month = f"{m.group(1)}-{m.group(2)}"
    return (estab.group(0) if estab else None), month


def parse_ecr_text(text: str) -> tuple[list[dict], list[str]]:
    rows, problems = [], []
    for n, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            continue
        parts = [p.strip() for p in line.split("#~#")]
        if len(parts) != ECR_FIELD_COUNT:
            problems.append(f"Line {n}: expected {ECR_FIELD_COUNT} fields, found {len(parts)}")
            continue
        uan, name, *nums = parts
        try:
            values = [int(float(v or 0)) for v in nums]
        except ValueError:
            problems.append(f"Line {n}: an amount is not a number")
            continue
        if any(v < 0 for v in values):
            problems.append(f"Line {n}: negative amount")
            continue
        gross, epf, eps, edli, ee, eps_c, diff, ncp, refund = values
        rows.append({
            "line": n, "uan": uan, "name": name, "gross": gross,
            "epf_wages": epf, "eps_wages": eps, "edli_wages": edli,
            "epf_contri": ee, "eps_contri": eps_c, "diff": diff,
            "ncp": ncp, "refund": refund,
        })
    return rows, problems


def _flag_combos(split: bool) -> list[tuple[bool, str, str, bool]]:
    """(eps_member, sep_basis, higher, higher_pension), most usual first."""
    combos = []
    for eps in (True, False):
        for sep in (("cap", "noeps", "new") if split and eps else ("cap",)):
            for higher in ("N", "E", "Y"):
                for hp in ((False, True) if eps and higher == "Y" else (False,)):
                    combos.append((eps, sep, higher, hp))
    return combos


def _wage_candidates(line: dict, wage_month: str) -> list[int]:
    """Full-month EPF wages worth trying, most natural first."""
    e, gross = line["epf_wages"], line["gross"]
    cands = [e]
    if wage_month == SPLIT_MONTH:
        f1 = SPLIT_DAYS_BEFORE / SPLIT_DAYS_TOTAL
        f2 = SPLIT_DAYS_AFTER / SPLIT_DAYS_TOTAL
        # The gross wage is often the EPF wage; below the new ceiling it is
        # the only value that splits back to exactly the ECR figure.
        if gross <= NEW_CEILING:
            cands.append(gross)
        for inv in ((e - OLD_CEILING * f1) / f2, e / f2):  # 1-30 Sep / 17-30 Sep
            base = math.floor(inv)
            cands += [base, base + 1, base - 1, base + 2, base - 2]
    cands.append(gross)
    seen, out = set(), []
    for c in cands:
        if c >= 0 and c not in seen:
            seen.add(c)
            out.append(c)
    return out


def infer_ecr_row(line: dict, wage_month: str) -> dict:
    """Register row (inputs + options) that reproduces one ECR line."""
    target = {k: line[k] for k in ("epf_wages", "eps_wages", "edli_wages", "eps_contri", "diff")}
    combos = _flag_combos(wage_month == SPLIT_MONTH)
    cands = _wage_candidates(line, wage_month)
    found = None
    # First pass: no VPF. Second: allow VPF to explain a higher employee share.
    for allow_vpf in (False, True):
        for w in cands:
            for eps, sep, higher, hp in combos:
                c = calc_row(w, wage_month, eps, sep, higher, 0, hp)
                if any(c[k] != v for k, v in target.items()):
                    continue
                vpf = line["epf_contri"] - c["epf_contri"]
                if vpf == 0 or (allow_vpf and vpf > 0):
                    found = (w, eps, sep, higher, hp, vpf)
                    break
            if found:
                break
        if found:
            break

    note = ""
    if found:
        w, eps, sep, higher, hp, vpf = found
    else:
        # Keep the file's EPF wage; say which figures this tool would change.
        w, sep, higher, hp, vpf = line["epf_wages"], "cap", "N", False, 0
        eps = line["eps_wages"] > 0 or line["epf_wages"] == 0
        c = calc_row(w, wage_month, eps, sep, higher, 0, hp)
        labels = {"epf_wages": "EPF wages", "eps_wages": "EPS wages",
                  "edli_wages": "EDLI wages", "epf_contri": "EPF contri.",
                  "eps_contri": "EPS contri.", "diff": "Diff"}
        diffs = [f"{label} {line[k]}→{c[k]}" for k, label in labels.items() if c[k] != line[k]]
        note = "Recalculated: " + ", ".join(diffs)

    return {
        "uan": re.sub(r"\D", "", line["uan"])[:12],
        "name": line["name"],
        "gross": line["gross"],
        "epf": w,
        "ncp": line["ncp"],
        "vpf": vpf,
        "refund": line["refund"],
        "eps": "Y" if eps else "N",
        "sep": sep,
        "higher": higher,
        "hp": "Y" if hp else "N",
        "note": note,
    }


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    _log("view")
    visitor = visitors.get_registration(session.get("reg_id"))
    return render_template("index.html", registered=visitor is not None, visitor=visitor)


@app.route("/api/upload", methods=["POST"])
def upload():
    """
    First call (no mapping): try to auto-detect columns; if all required
    fields are found, parse + return rows straight away. If not, return
    the headers + a guessed mapping so the browser can show a mapping
    step, then re-POST with `mapping` filled in (same file, resent).
    """
    import pandas as pd

    file = request.files.get("file")
    if file is None:
        return jsonify({"error": "No file uploaded"}), 400

    filename = file.filename or ""
    try:
        if filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
            df = pd.read_excel(file, dtype=str)
        else:
            df = pd.read_csv(file, dtype=str)
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": f"Could not read file: {exc}"}), 400

    df = df.dropna(how="all")
    headers = [str(c) for c in df.columns]

    mapping_json = request.form.get("mapping")
    if mapping_json:
        import json

        mapping = json.loads(mapping_json)
    else:
        mapping = guess_mapping(headers)
        missing_required = [f for f in REQUIRED_FIELDS if not mapping.get(f)]
        if missing_required:
            sample_rows = df.head(5).fillna("").to_dict(orient="records")
            return jsonify(
                {
                    "status": "needs_mapping",
                    "headers": headers,
                    "guessed_mapping": mapping,
                    "sample_rows": sample_rows,
                    "required_fields": REQUIRED_FIELDS,
                    "optional_fields": OPTIONAL_FIELDS,
                }
            )

    rows = []
    issues_by_row = []
    for _, r in df.iterrows():
        def get(field):
            col = mapping.get(field)
            if not col or col not in df.columns:
                return ""
            val = r.get(col, "")
            return "" if pd.isna(val) else str(val).strip()

        row = {
            "uan": re.sub(r"\D", "", get("uan"))[:12],
            "name": get("name"),
            "gross": get("gross") or "0",
            "epf": get("epf") or "0",
            "ncp": _whole_number_text(get("ncp")) or "0",
            "vpf": _whole_number_text(get("vpf")) or "0",
            "refund": get("refund") or "0",
        }
        rows.append(row)
        issues_by_row.append(validate_wage_row(row))

    _log("upload")
    return jsonify(
        {
            "status": "ok",
            "mapping": mapping,
            "rows": rows,
            "issues": issues_by_row,
        }
    )


@app.route("/api/import-ecr", methods=["POST"])
def import_ecr():
    """Read an ECR text file back into the wage register."""
    file = request.files.get("file")
    if file is None:
        return jsonify({"error": "No file uploaded"}), 400
    raw = file.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    if "#~#" not in text:
        return jsonify({"error": "This is not an ECR text file (no #~# separators found)"}), 400

    estab_code, file_month = ecr_filename_hints(file.filename)
    page_month = request.form.get("wage_month", "")
    wage_month = file_month or page_month
    if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", wage_month or ""):
        return jsonify({"error": "Choose the wage month before importing"}), 400

    lines, problems = parse_ecr_text(text)
    if not lines:
        return jsonify({"error": problems[0] if problems else "The file has no ECR lines"}), 400

    rows = [infer_ecr_row(line, wage_month) for line in lines]
    _log("import")
    return jsonify({
        "status": "ok",
        "rows": rows,
        "skipped": problems,
        "wage_month": wage_month,
        "month_from_file": bool(file_month),
        "estab_code": estab_code,
    })


# Header names match FIELD_SYNONYMS, so a filled-in template uploads without
# the mapping step.
TEMPLATE_COLUMNS = [
    ("UAN", 16),
    ("Member name", 30),
    ("Gross wages", 14),
    ("EPF wages", 14),
    ("NCP days", 11),
    ("VPF", 11),
    ("Refund of advances", 20),
]


TEMPLATE_FILENAME = "ECR_wage_upload_template.xlsx"


# The URL ends in the filename so browsers that name downloads from the
# URL (ignoring Content-Disposition) still save it as .xlsx.
@app.route("/api/template")
@app.route(f"/api/template/{TEMPLATE_FILENAME}")
def upload_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Wages"
    ws.append([name for name, _ in TEMPLATE_COLUMNS])
    for idx, (_, width) in enumerate(TEMPLATE_COLUMNS, start=1):
        col = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[col].width = width
        ws.cell(row=1, column=idx).font = Font(bold=True)
    # Store UANs as text so Excel doesn't turn 12 digits into 1.01E+11.
    for row in range(2, 1001):
        ws.cell(row=row, column=1).number_format = "@"
    ws.freeze_panes = "A2"

    notes = wb.create_sheet("Instructions")
    for line in [
        ["Fill one row per member on the Wages sheet, then upload this file."],
        ["UAN", "12-digit UAN (required)"],
        ["Member name", "As in the UAN record (required)"],
        ["Gross wages", "Total wages for the month, in rupees (required)"],
        ["EPF wages", "Full-month EPF wages, in rupees (required)"],
        ["NCP days", "Non-contributory period days: whole days, 0 to days in the month"],
        ["VPF", "Voluntary PF for the month, in rupees (0 if none)"],
        ["Refund of advances", "In rupees (0 if none)"],
    ]:
        notes.append(line)
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 70
    notes["A1"].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=TEMPLATE_FILENAME,
    )


@app.route("/api/ecr/generate", methods=["POST"])
def generate_ecr():
    data = request.get_json(force=True)
    estab_code = (data.get("estab_code") or "").strip()[:15]
    wage_month = data.get("wage_month")  # "YYYY-MM"
    rows = data.get("rows") or []

    if not estab_code or not wage_month:
        return jsonify({"error": "Establishment code and wage month are required"}), 400

    year, month = wage_month.split("-")

    for n, row in enumerate(rows, start=1):
        issue = ncp_issue(row.get("ncp"), wage_month) or vpf_issue(row.get("vpf"))
        if issue:
            return jsonify({"error": f"Row {n}: {issue}"}), 400
    dup = duplicate_uan_issue(rows)
    if dup:
        return jsonify({"error": dup}), 400

    lines = []
    d = "#~#"
    for row in rows:
        gross = float(row.get("gross") or 0)
        epf = float(row.get("epf") or 0)
        calc = calc_row(
            epf,
            wage_month,
            row_eps_member(row),
            row_sep_basis(row),
            row_higher_epf(row),
            float(row.get("vpf") or 0),
            row_higher_pension(row),
        )
        ncp = int(row.get("ncp") or 0)
        refund = row.get("refund") or 0

        dd = [
            str(row.get("uan", ""))[:12],
            str(row.get("name", "")).upper(),
            str(int(gross))[:10],
            str(calc["epf_wages"])[:10],
            str(calc["eps_wages"])[:10],
            str(calc["edli_wages"])[:10],
            str(calc["epf_contri"])[:10],
            str(calc["eps_contri"])[:10],
            str(calc["diff"])[:10],
            str(ncp),
            str(refund)[:10] if float(refund or 0) > 0 else "0",
        ]
        lines.append(d.join(dd))

    content = "\r\n".join(lines) + ("\r\n" if lines else "")
    buf = io.BytesIO(content.encode("utf-8"))
    buf.seek(0)
    filename = f"{estab_code}{year}{month}.txt"
    _log("ecr")
    return send_file(
        buf,
        mimetype="text/plain",
        as_attachment=True,
        download_name=filename,
    )


@app.route("/api/exit/generate", methods=["POST"])
def generate_exit():
    data = request.get_json(force=True)
    estab_code = (data.get("estab_code") or "").strip()[:15]
    wage_month = data.get("wage_month")
    rows = data.get("rows") or []

    if not estab_code or not wage_month:
        return jsonify({"error": "Establishment code and wage month are required"}), 400

    year, month = wage_month.split("-")

    lines = []
    d = "#~#"
    for row in rows:
        uan = str(row.get("uan", ""))[:12]
        name = str(row.get("name", ""))
        raw_date = row.get("date", "")
        code = str(row.get("code", ""))[:1]

        date_str = ""
        if raw_date:
            try:
                dt = datetime.strptime(raw_date, "%Y-%m-%d")
                date_str = dt.strftime("%d-%m-%Y")
            except ValueError:
                date_str = raw_date

        dd = [uan, date_str, code]
        lines.append(d.join(dd))

    content = "\r\n".join(lines) + ("\r\n" if lines else "")
    buf = io.BytesIO(content.encode("utf-8"))
    buf.seek(0)
    filename = f"{estab_code}{year}{month}exit.txt"
    _log("exit")
    return send_file(
        buf,
        mimetype="text/plain",
        as_attachment=True,
        download_name=filename,
    )


# ---------------------------------------------------------------------------
# Arrear file (EPFO Unified Portal: ECR / Returns > Arrear File Upload)
#
# 8 #~#-separated fields per member: UAN, member name, arrear EPF wages,
# arrear EPS wages, arrear EDLI wages, EPF (employee) contribution due,
# EPS contribution due, ER (employer EPF diff) contribution due.
# The arrear period (from / till month), disbursement date, contribution rate
# and remarks are entered on the portal's upload form, so one file covers one
# period at one rate.
#
# Each member's arrear is worked out month by month from the EPF wage already
# paid and the revised EPF wage. EPS and EDLI arrears (and capped EPF, when
# higher EPF is "No") only count up to that month's wage ceiling: 15,000 up
# to August 2026, 25,000 from October 2026, and September 2026 split 16/30 at
# 15,000 and 14/30 at 25,000, as in the wage register.
# Keep calc_arrear() in sync with calcArrear() in static/app.js.
# ---------------------------------------------------------------------------
ARREAR_RATES = {"12": 0.12, "10": 0.10}
MONTH_RE = re.compile(r"20\d{2}-(0[1-9]|1[0-2])")
MAX_ARREAR_MONTHS = 120


def months_between(from_month: str, till_month: str) -> list[str]:
    """Months from `from_month` to `till_month` inclusive (YYYY-MM)."""
    if not (MONTH_RE.fullmatch(from_month or "") and MONTH_RE.fullmatch(till_month or "")):
        return []
    y, m = (int(p) for p in from_month.split("-"))
    ty, tm = (int(p) for p in till_month.split("-"))
    out = []
    while (y, m) <= (ty, tm) and len(out) <= MAX_ARREAR_MONTHS:
        out.append(f"{y}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def capped_month_wage(w: float, month: str) -> float:
    """The part of a monthly wage that falls within that month's ceiling."""
    if month == SPLIT_MONTH:
        f1 = SPLIT_DAYS_BEFORE / SPLIT_DAYS_TOTAL
        f2 = SPLIT_DAYS_AFTER / SPLIT_DAYS_TOTAL
        return min(w, OLD_CEILING) * f1 + min(w, NEW_CEILING) * f2
    return min(w, ceiling_for(month))


def _override(value):
    s = str(value if value is not None else "").strip()
    return int(float(s)) if re.fullmatch(r"\d+(\.\d+)?", s) else None


def calc_arrear(paid, revised, months, eps_member=True, higher="N", rate="12",
                override: dict | None = None) -> dict:
    """
    `paid` and `revised` are the member's monthly EPF wage before and after
    the revision; `months` the arrear period. `higher` is a HIGHER_MODES
    value: Y = employee and employer on the full arrear, E = employee on the
    full arrear and employer only up to the ceiling, N = both up to the
    ceiling. `override` may set "epf", "eps" or "edli" arrear wages by hand.
    """
    paid, revised = float(paid or 0), float(revised or 0)
    mode = higher if higher in HIGHER_MODES else "N"
    full = (revised - paid) * len(months)
    capped = sum(capped_month_wage(revised, m) - capped_month_wage(paid, m) for m in months)
    wages = {
        "epf": round_half_up(capped if mode == "N" else full),
        "eps": round_half_up(capped) if eps_member else 0,
        "edli": round_half_up(capped),
    }
    for key, value in (override or {}).items():
        value = _override(value)
        if key in wages and value is not None:
            wages[key] = value
    # Employer's share follows the EPF arrear, except "employee only" where
    # it stays within the ceiling.
    employer_w = min(wages["epf"], round_half_up(capped)) if mode == "E" else wages["epf"]
    r = ARREAR_RATES.get(str(rate), 0.12)
    ee = round_half_up(wages["epf"] * r)
    eps_c = round_half_up(wages["eps"] * 0.0833)
    return {
        "epf_wages": wages["epf"],
        "eps_wages": wages["eps"],
        "edli_wages": wages["edli"],
        "ee": ee,
        "eps": eps_c,
        "er": round_half_up(employer_w * r) - eps_c,
    }


def arrear_row_issue(row: dict) -> str:
    try:
        paid = float(row.get("paid") or 0)
        revised = float(row.get("revised") or 0)
    except (TypeError, ValueError):
        return "wages must be numbers"
    if paid < 0 or revised < 0:
        return "wages cannot be negative"
    if revised < paid:
        return "revised wage is less than the wage paid"
    for key in ("epf", "eps", "edli"):
        value = (row.get("ovr") or {}).get(key)
        if value not in (None, "") and _override(value) is None:
            return f"arrear {key.upper()} wages must be a whole number"
    return ""


def _arrear_filename(estab_code: str, from_month: str, till_month: str) -> str:
    # Portal rule: letters and digits only, at most 32 characters.
    code = re.sub(r"[^A-Za-z0-9]", "", estab_code)[:15]
    return f"{code}ARR{from_month.replace('-', '')}{till_month.replace('-', '')}.txt"


@app.route("/api/arrear/generate", methods=["POST"])
def generate_arrear():
    data = request.get_json(force=True)
    estab_code = (data.get("estab_code") or "").strip()
    from_month, till_month = data.get("from_month", ""), data.get("till_month", "")
    rate = str(data.get("rate", "12"))
    rows = data.get("rows") or []

    if not estab_code:
        return jsonify({"error": "Establishment code is required"}), 400
    months = months_between(from_month, till_month)
    if not months:
        return jsonify({"error": "Choose the arrear From and Till months (From must not be after Till)"}), 400
    if len(months) > MAX_ARREAR_MONTHS:
        return jsonify({"error": f"The arrear period is longer than {MAX_ARREAR_MONTHS} months"}), 400
    if rate not in ARREAR_RATES:
        return jsonify({"error": "Contribution rate must be 12% or 10%"}), 400
    if not rows:
        return jsonify({"error": "Add at least one member to the arrear register"}), 400
    for n, row in enumerate(rows, start=1):
        if not re.fullmatch(r"\d{12}", str(row.get("uan", "")).strip()):
            return jsonify({"error": f"Row {n}: UAN is not a 12-digit number"}), 400
        issue = arrear_row_issue(row)
        if issue:
            return jsonify({"error": f"Row {n}: {issue}"}), 400
    dup = duplicate_uan_issue(rows)
    if dup:
        return jsonify({"error": dup}), 400

    lines = []
    for row in rows:
        c = calc_arrear(row.get("paid"), row.get("revised"), months, row_eps_member(row),
                        row_higher_epf(row), rate, row.get("ovr"))
        lines.append("#~#".join([
            str(row.get("uan", "")).strip()[:12],
            str(row.get("name", "")).strip().upper(),
            str(c["epf_wages"]), str(c["eps_wages"]), str(c["edli_wages"]),
            str(c["ee"]), str(c["eps"]), str(c["er"]),
        ]))

    content = "\r\n".join(lines) + "\r\n"
    _log("arrear")
    return send_file(
        io.BytesIO(content.encode("utf-8")),
        mimetype="text/plain",
        as_attachment=True,
        download_name=_arrear_filename(estab_code, from_month, till_month),
    )


ARREAR_FIELD_SYNONYMS = {
    "uan": ["uan", "universal account number", "member uan"],
    "name": ["member name", "employee name", "name"],
    "paid": ["epf wage paid", "epf wages paid", "wage paid", "old epf wage", "existing epf wage"],
    "revised": ["revised epf wage", "revised epf wages", "revised wage", "new epf wage"],
    "eps": ["eps member"],
    "higher": ["higher epf"],
    "arr_epf": ["arrear epf wages", "arrear epf wage"],
    "arr_eps": ["arrear eps wages", "arrear eps wage"],
    "arr_edli": ["arrear edli wages", "arrear edli wage"],
}
ARREAR_REQUIRED = ["uan", "name", "paid", "revised"]

ARREAR_TEMPLATE_FILENAME = "ECR_arrear_upload_template.xlsx"
ARREAR_TEMPLATE_COLUMNS = [
    ("UAN", 16),
    ("Member name", 28),
    ("EPF wage paid (per month)", 24),
    ("Revised EPF wage (per month)", 26),
    ("EPS member (Y/N)", 16),
    ("Higher EPF (Y/E/N)", 17),
    ("Arrear EPF wages (optional)", 25),
    ("Arrear EPS wages (optional)", 25),
    ("Arrear EDLI wages (optional)", 26),
]
ARREAR_TEMPLATE_SAMPLES = [
    ["100000000001", "SAMPLE MEMBER ONE", 12000, 14000, "Y", "N", None, None, None],
    ["100000000002", "SAMPLE MEMBER TWO", 14000, 17000, "Y", "N", None, None, None],
    ["100000000003", "SAMPLE MEMBER THREE", 20000, 22000, "N", "Y", None, None, None],
]


def _guess_columns(headers: list[str], synonyms: dict[str, list[str]]) -> dict[str, str | None]:
    normalized = {h: _normalize_header(h) for h in headers}
    mapping: dict[str, str | None] = {}
    taken: set[str] = set()
    for exact in (True, False):
        for field, words in synonyms.items():
            if mapping.get(field):
                continue
            for h, norm in normalized.items():
                if h in taken:
                    continue
                if (norm in words) if exact else any(norm.startswith(w) for w in words):
                    mapping[field] = h
                    taken.add(h)
                    break
    return {f: mapping.get(f) for f in synonyms}


@app.route("/api/arrear/upload", methods=["POST"])
def upload_arrear():
    import pandas as pd

    file = request.files.get("file")
    if file is None:
        return jsonify({"error": "No file uploaded"}), 400
    filename = (file.filename or "").lower()
    try:
        if filename.endswith((".xlsx", ".xlsm", ".xls")):
            df = pd.read_excel(file, dtype=str)
        else:
            df = pd.read_csv(file, dtype=str)
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": f"Could not read file: {exc}"}), 400
    df = df.dropna(how="all")

    mapping = _guess_columns([str(c) for c in df.columns], ARREAR_FIELD_SYNONYMS)
    missing = [f for f in ARREAR_REQUIRED if not mapping.get(f)]
    if missing:
        names = {"uan": "UAN", "name": "Member name", "paid": "EPF wage paid (per month)",
                 "revised": "Revised EPF wage (per month)"}
        return jsonify({"error": "Columns not found: " + ", ".join(names[f] for f in missing)
                        + ". Use the arrear sample Excel."}), 400

    rows = []
    for _, r in df.iterrows():
        def get(field):
            col = mapping.get(field)
            val = r.get(col, "") if col else ""
            return "" if pd.isna(val) else _whole_number_text(str(val).strip())

        eps = get("eps").upper()[:1]
        higher = get("higher").upper()[:1]
        rows.append({
            "uan": re.sub(r"\D", "", get("uan"))[:12],
            "name": get("name"),
            "paid": get("paid") or "0",
            "revised": get("revised") or "0",
            "eps": "N" if eps == "N" else "Y",
            "higher": higher if higher in HIGHER_MODES else "N",
            "ovr": {k: get(f"arr_{k}") for k in ("epf", "eps", "edli") if get(f"arr_{k}")},
        })
    _log("upload")
    return jsonify({"status": "ok", "rows": rows})


@app.route("/api/arrear/template")
@app.route(f"/api/arrear/template/{ARREAR_TEMPLATE_FILENAME}")
def arrear_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Arrears"
    ws.append([name for name, _ in ARREAR_TEMPLATE_COLUMNS])
    for idx, (_, width) in enumerate(ARREAR_TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        ws.column_dimensions[cell.column_letter].width = width
        cell.font = Font(bold=True)
    for sample in ARREAR_TEMPLATE_SAMPLES:
        ws.append(sample)
    # Store UANs as text so Excel doesn't turn 12 digits into 1.01E+11.
    for row in range(2, 1001):
        ws.cell(row=row, column=1).number_format = "@"
    ws.freeze_panes = "A2"

    notes = wb.create_sheet("Instructions")
    for line in [
        ["Arrear upload sheet — one row per member. Replace the three sample rows with your members."],
        [],
        ["UAN", "12-digit UAN (required)"],
        ["Member name", "As in the UAN record (required)"],
        ["EPF wage paid (per month)", "Monthly EPF wage (Basic + DA) on which PF was already paid for the arrear period (required)"],
        ["Revised EPF wage (per month)", "Monthly EPF wage after the revision (required)"],
        ["EPS member (Y/N)", "N if the member is not in the pension scheme (EPS arrear wages = 0). Blank = Y"],
        ["Higher EPF (Y/E/N)", "Y = both shares on the full arrear; E = employee on the full arrear, employer up to the ceiling; N or blank = both up to the ceiling"],
        ["Arrear EPF / EPS / EDLI wages", "Optional. Leave blank to let the generator calculate them; fill in only to override"],
        [],
        ["How it is calculated", "Choose the From and Till months (the arrear period) and the contribution rate on the Arrear Register tab."],
        ["", "Arrear EPF wages = (revised wage - wage paid) x number of months; with Higher EPF = N, only the part within each month's ceiling."],
        ["", "Arrear EPS and EDLI wages = only the part within each month's ceiling: Rs 15,000 up to Aug 2026, Rs 25,000 from Oct 2026, Sep 2026 split 16/30 and 14/30."],
        ["", "Employee share = rate x arrear EPF wages; EPS = 8.33% of arrear EPS wages; employer EPF (diff) = rate x arrear EPF wages - EPS."],
        [],
        ["Samples (arrear period Apr 2026 to Jun 2026, 3 months, 12%)", ""],
        ["SAMPLE MEMBER ONE", "12,000 -> 14,000: arrear EPF/EPS/EDLI wages 6,000; EE 720, EPS 500, ER 220"],
        ["SAMPLE MEMBER TWO", "14,000 -> 17,000: only 1,000 a month is within the 15,000 ceiling, so arrear wages 3,000; EE 360, EPS 250, ER 110"],
        ["SAMPLE MEMBER THREE", "20,000 -> 22,000, not in EPS, higher EPF on both: arrear EPF 6,000, EPS 0, EDLI 0; EE 720, EPS 0, ER 720"],
    ]:
        notes.append(line)
    notes.column_dimensions["A"].width = 34
    notes.column_dimensions["B"].width = 120
    notes["A1"].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=ARREAR_TEMPLATE_FILENAME,
    )


# ---------------------------------------------------------------------------
# POHW arrear file (EPFO Unified Portal: POHW Arrear File Upload) —
# pension on higher wages. 8 #~#-separated fields per member: UAN, member
# name, arrear EPF wages, arrear EPS wages, arrear EDLI wages, arrear EPF EE
# share, arrear EPF ER share, arrear EPS share. Per the POHW help file all
# wages and shares are 0 except the EPS share, which must be more than 0:
# the amount EPFO has worked out for the member's pension on higher wages.
# The wage month (which must be the current wage month), contribution rate
# and remarks are entered on the portal form; one POHW arrear per UAN.
# ---------------------------------------------------------------------------
POHW_TEMPLATE_FILENAME = "ECR_POHW_arrear_upload_template.xlsx"
POHW_FIELD_SYNONYMS = {
    "uan": ["uan", "universal account number", "member uan"],
    "name": ["member name", "employee name", "name"],
    "eps_share": ["arrear eps share", "eps share", "pohw amount", "eps amount"],
}


def pohw_amount_issue(value) -> str:
    s = str(value if value is not None else "").strip().replace(",", "")
    if not re.fullmatch(r"\d+(\.0+)?", s):
        return "EPS share must be a whole number of rupees"
    return "" if int(float(s)) > 0 else "EPS share must be more than 0"


def _pohw_filename(estab_code: str, wage_month: str) -> str:
    # Portal rule: letters and digits only, at most 32 characters.
    code = re.sub(r"[^A-Za-z0-9]", "", estab_code)[:15]
    return f"{code}POHW{wage_month.replace('-', '')}.txt"


@app.route("/api/pohw/generate", methods=["POST"])
def generate_pohw():
    data = request.get_json(force=True)
    estab_code = (data.get("estab_code") or "").strip()
    wage_month = data.get("wage_month", "")
    rows = data.get("rows") or []
    if not estab_code:
        return jsonify({"error": "Establishment code is required"}), 400
    if not MONTH_RE.fullmatch(wage_month or ""):
        return jsonify({"error": "Choose the wage month"}), 400
    if not rows:
        return jsonify({"error": "Add at least one member to the POHW arrear register"}), 400
    for n, row in enumerate(rows, start=1):
        if not re.fullmatch(r"\d{12}", str(row.get("uan", "")).strip()):
            return jsonify({"error": f"Row {n}: UAN is not a 12-digit number"}), 400
        issue = pohw_amount_issue(row.get("eps_share"))
        if issue:
            return jsonify({"error": f"Row {n}: {issue}"}), 400
    dup = duplicate_uan_issue(rows)
    if dup:
        return jsonify({"error": dup}), 400

    lines = []
    for row in rows:
        amount = int(float(str(row.get("eps_share")).strip().replace(",", "")))
        lines.append("#~#".join([
            str(row.get("uan", "")).strip()[:12],
            str(row.get("name", "")).strip().upper(),
            "0", "0", "0", "0", "0", str(amount),
        ]))
    content = "\r\n".join(lines) + "\r\n"
    _log("pohw")
    return send_file(
        io.BytesIO(content.encode("utf-8")),
        mimetype="text/plain",
        as_attachment=True,
        download_name=_pohw_filename(estab_code, wage_month),
    )


@app.route("/api/pohw/upload", methods=["POST"])
def upload_pohw():
    import pandas as pd

    file = request.files.get("file")
    if file is None:
        return jsonify({"error": "No file uploaded"}), 400
    filename = (file.filename or "").lower()
    try:
        if filename.endswith((".xlsx", ".xlsm", ".xls")):
            df = pd.read_excel(file, dtype=str)
        else:
            df = pd.read_csv(file, dtype=str)
    except Exception as exc:  # noqa: BLE001
        return jsonify({"error": f"Could not read file: {exc}"}), 400
    df = df.dropna(how="all")

    mapping = _guess_columns([str(c) for c in df.columns], POHW_FIELD_SYNONYMS)
    missing = [f for f in POHW_FIELD_SYNONYMS if not mapping.get(f)]
    if missing:
        names = {"uan": "UAN", "name": "Member name", "eps_share": "Arrear EPS share"}
        return jsonify({"error": "Columns not found: " + ", ".join(names[f] for f in missing)
                        + ". Use the POHW sample Excel."}), 400

    rows = []
    for _, r in df.iterrows():
        def get(field):
            val = r.get(mapping[field], "")
            return "" if pd.isna(val) else _whole_number_text(str(val).strip())

        rows.append({
            "uan": re.sub(r"\D", "", get("uan"))[:12],
            "name": get("name"),
            "eps_share": get("eps_share").replace(",", ""),
        })
    _log("upload")
    return jsonify({"status": "ok", "rows": rows})


@app.route("/api/pohw/template")
@app.route(f"/api/pohw/template/{POHW_TEMPLATE_FILENAME}")
def pohw_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "POHW arrears"
    for idx, (name, width) in enumerate(
        [("UAN", 16), ("Member name", 30), ("Arrear EPS share", 18)], start=1
    ):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.append(["100000000001", "SAMPLE MEMBER ONE", 15000])
    ws.append(["100000000002", "SAMPLE MEMBER TWO", 48250])
    for row in range(2, 1001):
        ws.cell(row=row, column=1).number_format = "@"
    ws.freeze_panes = "A2"

    notes = wb.create_sheet("Instructions")
    for line in [
        ["POHW (pension on higher wages) arrear upload sheet — one row per member. Replace the two sample rows."],
        [],
        ["UAN", "12-digit UAN (required)"],
        ["Member name", "As in the UAN record (required)"],
        ["Arrear EPS share", "Amount in whole rupees to be remitted to EPS for the member's pension on higher wages, as worked out by EPFO (required, more than 0)"],
        [],
        ["File format", "UAN#~#NAME#~#0#~#0#~#0#~#0#~#0#~#EPS share — arrear EPF, EPS and EDLI wages and the EE and ER shares are always 0"],
        ["On the portal", "POHW Arrear File Upload: the wage month must be the current wage month; choose the contribution rate and enter remarks there"],
        ["", "A UAN that already has an approved, in-process or paid POHW arrear ECR cannot be filed again"],
    ]:
        notes.append(line)
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 120
    notes["A1"].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=POHW_TEMPLATE_FILENAME,
    )


@app.route("/healthz")
def healthz():
    return jsonify({"status": "ok"})


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
